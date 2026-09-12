from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Kait_当前版本游戏设计文档_v0.8.2.docx"
SCREENSHOT = ROOT / "Logs" / "drag-target-direction.png"

NAVY = "243044"
BLUE = "2E6FBA"
ICE = "DDEEFF"
CREAM = "FFF6DF"
GOLD = "D7A742"
SILVER = "AEB8C6"
INK = "262A33"
MUTED = "667085"
WHITE = "FFFFFF"
RED = "B6424E"


KAIT_ACTIVE = [
    ("A01", "大步奔行之靴", "普通", "冷却2", "本回合速度 +1。用于修正路线、补足连杀距离。"),
    ("A02", "咒剑诅咒", "普通", "冷却3", "指定敌人下一次受到的伤害 +1。"),
    ("A03", "怪物定身术", "普通", "冷却3", "指定敌人跳过下一次行动。"),
    ("A04", "解除魔法", "普通", "冷却4", "移除一个裂隙，腾出安全空间。"),
    ("A05", "雷鸣波", "罕见", "冷却5", "下一次移动将同方向上的全部敌人推到尽头；Kait 本人不移动，威胁盘仍响应。"),
    ("A06", "次级幻影", "罕见", "冷却4", "指定敌人成为敌方下一阶段的共同攻击目标。"),
    ("A07", "命令术", "罕见", "冷却3", "将一个已锁定的攻击方向顺时针旋转 90°。"),
    ("A08", "迷踪步", "罕见", "冷却3", "移动到相邻空格，不推动威胁盘。"),
    ("A09", "哈达之握", "罕见", "冷却4", "将同行或同列的敌人拉到 Kait 面前。"),
    ("A10", "猫之迅捷", "稀有", "冷却5", "本回合速度翻倍，并在固定速度加成之前计算。"),
    ("A11", "魔能斩", "稀有", "冷却4", "下一次未击杀命中把目标推到路径尽头。"),
    ("A12", "不息咒缚", "稀有", "冷却4", "传送到相邻的受诅咒敌人附近，形成追击位置。"),
    ("A13", "列维斯图斯之墓", "稀有", "冷却5", "本回合不能移动，但也不会受到伤害。"),
    ("A14", "任意门", "稀有", "冷却4", "下一次裂隙改在棋盘中心对称位置出现。"),
]

KAIT_PASSIVE = [
    ("P01", "警戒武器", "普通", "—", "额外显示接下来的 2 个威胁数字。"),
    ("P02", "非致命击倒", "普通", "—", "敌方友伤不会把目标生命降到 1 以下。"),
    ("P03", "震慑斩", "普通", "—", "敌人被推撞柱子后跳过下一次行动。"),
    ("P04", "秘法锁", "普通", "—", "被占用的裂隙可延迟一次刷新。"),
    ("P05", "衰弱射线", "普通", "—", "受到友伤的敌人跳过下一次行动。"),
    ("P06", "守契者权杖", "罕见", "—", "每回合第一次使用主动技能时，另一张随机主动技能冷却 -1。"),
    ("P07", "刃之契约", "罕见", "—", "一次连杀首次达到 3 杀时，使当前冷却最长的技能冷却 -1。"),
    ("P08", "咒术护甲", "罕见", "—", "受诅咒敌人的下一次攻击必定失败。"),
    ("P09", "咒术大师", "罕见", "—", "受诅咒敌人死亡后，诅咒转移到本次连杀的下一个目标。"),
    ("P10", "疯狂咒缚", "罕见", "—", "每回合第一次诅咒伤害对目标相邻敌人各造成 1 点伤害。"),
    ("P11", "幸运剑", "罕见", "—", "每次奖励可免费重抽一次。"),
    ("P12", "饮命者", "罕见", "—", "每次连杀第一次击杀受诅咒敌人时恢复 1 点生命。"),
    ("P13", "守卫刻文", "罕见", "—", "连杀结束时留下刻文；下一次被阻挡的裂隙会转移到刻文处。"),
    ("P14", "念动力", "稀有", "—", "每回合第一次推动敌人时，使威胁盘对应数字移动 1 格，且不合并。"),
    ("P15", "拟像术", "稀有", "—", "复制另一张被动技能，但自身仍占用一个槽位。"),
    ("P16", "诅咒幽魂", "稀有", "—", "每次连杀第一次击杀会留下幽魂，阻挡下一次攻击。"),
    ("P17", "斥力魔爆", "稀有", "—", "推动会向后额外传递给一个敌人。"),
    ("P18", "鸦后预兆", "稀有", "—", "新生成的 2 优先出现在实际威胁移动方向的反侧，无法放置时回退。"),
    ("P19", "移位斗篷", "稀有", "—", "每个敌方阶段第一次受到的远程攻击失效。"),
    ("P20", "远古奥秘之书", "稀有", "—", "当盘面至少有 5 个 2 时，每回合一次把最早的两个 2 合成为 4。"),
    ("P21", "穿墙术", "稀有", "—", "威胁数字经过柱子时不再被柱子截停。"),
    ("P22", "异次元袋", "稀有", "—", "合成后收纳来源后方一个同值数字，优先处理移动方向后侧。"),
    ("P23", "重力反转", "稀有", "—", "威胁盘始终按玩家输入的反方向移动。"),
    ("P24", "化零为整", "稀有", "—", "每回合第一次将两个同阶裂隙合成为更高阶裂隙。"),
    ("P25", "狂野魔法涌动", "稀有·实验", "—", "裂隙随机偏移，且新敌人跳过第一次行动。正式牌池默认关闭。"),
]

YUMMN_ACTIVE = [
    ("R02", "疾步如风·收势", "普通", "1气", "选择方向后恰好移动 1 格，总计只消耗 1 气，不出拳。"),
    ("E03", "塑造流水", "罕见", "2气", "在任意合法空格制造永久冰柱；场上只保留最新一根。"),
    ("E04", "冬之吐息", "罕见", "2气", "选择方向，前方最先覆盖的 2 格各造成 1 点伤害，存活目标被冻结。"),
    ("S01", "暗影步", "普通", "1气", "传送到指定方向最近的空残影位置。"),
    ("R17", "影遁术", "罕见", "2气", "传送到指定方向的空残影并消耗该残影，本次不补 2。"),
    ("S02", "黑暗术", "罕见", "2气", "在任意合法空格制造永久暗幕；阻挡箭矢，处于其中的敌人不能攻击；只保留最新一处。"),
    ("R26", "疗伤冥想", "罕见", "3气", "点选 Yummn 自身释放，恢复 1 点生命后推进敌方阶段。"),
    ("R30", "空震掌", "罕见", "2气", "对指定方向第一个敌人造成 1 点伤害，并将其推到障碍前。"),
    ("R38", "空冥身", "罕见", "3气", "沿指定方向穿过敌人，不攻击，停在路径最后一个空格。"),
    ("R39", "次级幻影", "罕见", "1气", "在任意合法空格放置唯一诱饵；敌人下次锁定它，承受一次攻击后消失。"),
]

YUMMN_PASSIVE = [
    ("R01", "疾风连击", "罕见", "—", "每次拳击改为消耗 2 气的二连击；反应攻击同样生效并继承逐拳增幅。气不足时退化为普通拳击。"),
    ("R03", "风行之靴", "稀有", "—", "高速移动固定消耗 1 气；有效的主动移动会补充两个 2。"),
    ("O02", "散打技巧·推掌", "普通", "—", "每次未击杀拳击将目标免费推动 1 格。"),
    ("R05", "震慑拳", "罕见", "—", "首次命中未震慑且存活目标时额外消耗 1 气并震慑其完整一个敌方阶段；持续期间不重复耗气或刷新。"),
    ("R06", "无甲防御", "普通", "—", "执行等待后，抵挡该敌方阶段第一次伤害；治疗式等待也计入。"),
    ("M04", "拨挡飞弹", "罕见", "—", "每次将受到箭矢伤害时可消耗 1 气抵消。"),
    ("R08", "水鞭", "罕见", "—", "方向上存在非相邻敌人时，原地造成 1 伤并将其拉到面前。"),
    ("R09", "不坏气拳", "罕见", "—", "每次拳击把存活目标推到障碍前，优先于普通推动。"),
    ("E05", "火蛇之牙", "罕见", "—", "每次拳击使主目标身后一格的敌人额外受到 1 点伤害。"),
    ("R13", "寒霜之触", "罕见", "—", "成功强制移动敌人后使其冻结。"),
    ("E06", "碎冰掌", "稀有", "—", "拳击冻结目标会破冰；目标免疫这次首击，周围敌人各受 1 伤。"),
    ("R15", "冰封拳", "稀有", "—", "未击杀拳击会冻结目标。"),
    ("R19", "暗影斗篷", "稀有", "—", "所有残影在一个敌方阶段内可以承受任意次数攻击。"),
    ("S06", "伺机而动", "罕见", "—", "敌人进入相邻攻击范围时自动反击；每个敌人每次输入最多触发一次。"),
    ("R40", "借机攻击", "罕见", "—", "敌人离开相邻攻击范围前自动拳击一次；每个敌人每次输入最多触发一次，玩家主动离开不触发。"),
    ("R22", "暗影反击", "罕见", "—", "残影受到攻击时向攻击者反射 1 点伤害；每次攻击结算一次。"),
    ("R23", "斗战冥想", "稀有", "—", "主动移动不再补 2；每次击杀改为总共补充两个 2。"),
    ("R24", "静谧心境", "罕见", "—", "主动移动不再补 2；每次等待补充一个 2。"),
    ("O05", "混元体", "稀有", "—", "从气竭恢复到高速状态时恢复 1 点生命。"),
    ("M05", "完美自我", "普通", "—", "首次进入气竭时立刻恢复 1 气，但仍需补满气槽才能退出气竭。"),
    ("R28", "气海扩张", "罕见", "—", "最大气值 +3，同时每次击杀回气 -1。"),
    ("O06", "震颤掌", "稀有", "—", "拳击在目标上留下命中方向标记；下一次从不同方向命中额外造成 2 伤并移除对应标记，标记数量不受限。"),
    ("R31", "穿墙术", "罕见", "—", "威胁数字可以穿过柱子，不在柱子前停下。"),
    ("R32", "异次元袋", "罕见", "—", "合成后移除一个同来源小数字，优先选择移动方向后侧。"),
    ("R33", "重力反转", "罕见", "—", "威胁盘按玩家输入的反方向移动。"),
    ("R34", "远古奥秘之书", "普通", "—", "盘面出现 5 个 2 时，将最早的两个 2 合成为 4。"),
    ("O03", "追身步", "普通", "—", "近身推动成功后跟进目标原位置；属于有效移动，可生成残影。"),
    ("R37", "灵体护身", "稀有", "—", "高速状态受伤时消耗 3 气抵挡；同时每次击杀回气 -1。"),
]

ENEMIES = [
    ("4", "杂兵", "2", "贴脸攻击 1；可推动。承担占位、跳板与友伤链条。"),
    ("8", "剑士", "3", "贴脸攻击 1；基础近战阻挡单位。"),
    ("16", "弓手", "2", "锁定一回合后沿整排射击 1；命中第一个单位停止，敌我皆伤。"),
    ("32", "重甲兵", "4", "贴脸攻击 1；可推动，耐久更高。"),
    ("64", "术士", "2", "锁定目标格一回合，下一回合对十字区域造成 1 伤，敌我皆伤。"),
    ("128", "盾骑士 Boss", "8", "每个敌方阶段面向目标并攻击整排；正面伤害为 0，侧背正常。Kait 与 Yummn 都以 128 作为当前 Boss 生成阈值。"),
]


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def no_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Kait v0.8.2  ·  ")
    set_font(run, 8, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def style_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.75)
    sec.right_margin = Cm(1.75)
    sec.header_distance = Cm(0.75)
    sec.footer_distance = Cm(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 14),
        ("Subtitle", 13, MUTED, 0, 16),
        ("Heading 1", 18, NAVY, 12, 7),
        ("Heading 2", 13, BLUE, 9, 5),
        ("Heading 3", 10.5, INK, 7, 3),
    ):
        s = doc.styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size)
        s.font.bold = name != "Subtitle"
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    for sec in doc.sections:
        p = sec.header.paragraphs[0]
        p.text = "KAIT  ·  当前版本游戏设计文档"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.runs[0], 8, True, NAVY)
        add_page_number(sec.footer.paragraphs[0])


def add_title(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Kait")
    set_font(r, 34, True, NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("当前版本游戏设计文档")
    set_font(r, 20, True, BLUE)

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("v0.8.2 · 双角色、双棋盘与完整技能牌池")
    set_font(r, 12, False, MUTED)

    if SCREENSHOT.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SCREENSHOT), width=Inches(6.2))
        cap = doc.add_paragraph("当前实现快照：左侧战术盘、中央角色信息与输入、右侧 2048 威胁盘")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cap.runs[0], 8.5, False, MUTED)

    box = doc.add_table(rows=1, cols=2)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    box.columns[0].width = Inches(1.5)
    box.columns[1].width = Inches(4.6)
    entries = [
        ("文档性质", "实现态设计说明 / 实习 Demo 交付"),
        ("基准版本", "v0.8.2-exit-stun"),
        ("更新时间", "2026 年 9 月 13 日"),
        ("内容范围", "玩法、时序、角色、敌人、设置、存档、表现与全部技能"),
    ]
    for i, (k, v) in enumerate(entries):
        if i:
            box.add_row()
        row = box.rows[i]
        set_cell_fill(row.cells[0], NAVY)
        set_cell_fill(row.cells[1], CREAM)
        for c in row.cells:
            set_cell_margins(c, 80, 110, 80, 110)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].text = k
        row.cells[1].text = v
        set_font(row.cells[0].paragraphs[0].runs[0], 9, True, WHITE)
        set_font(row.cells[1].paragraphs[0].runs[0], 9, False, INK)
    doc.add_page_break()


def add_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True, color=NAVY)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(item))


def add_callout(doc, title, text, fill=ICE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, 110, 135, 110, 135)
    p = cell.paragraphs[0]
    r = p.add_run(title + "  ")
    set_font(r, 9.5, True, NAVY)
    r = p.add_run(text)
    set_font(r, 9.2, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_simple_table(doc, headers, rows, widths=None, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_fill(cell, header_fill)
        set_cell_margins(cell)
        set_font(cell.paragraphs[0].runs[0], 8.2, True, WHITE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            cell.width = Inches(widths[i])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    no_row_split(table.rows[0])
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        no_row_split(row)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.text = str(value)
            if ridx % 2:
                set_cell_fill(cell, "F5F7FA")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                for run in p.runs:
                    set_font(run, 7.8, i == 1, INK)
                if i in (0, 2, 3):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_skill_section(doc, heading, intro, skills):
    doc.add_heading(heading, level=2)
    add_p(doc, intro)
    add_simple_table(
        doc,
        ["编号", "技能", "稀有度", "冷却/气耗", "当前效果"],
        skills,
        widths=[0.48, 1.22, 0.72, 0.78, 3.65],
    )


def build():
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("目录", level=1)
    toc = [
        "1. 项目定位与版本范围", "2. 核心体验与完整循环", "3. 双棋盘规则与全局时序",
        "4. 角色规则：Kait", "5. 角色规则：Yummn", "6. 敌人与 Boss",
        "7. 完整技能牌池", "8. 奖励、构筑与装备", "9. 操作与界面",
        "10. 视觉、动画与音频", "11. 设置、存档与平台", "12. 胜负、平衡与测试",
    ]
    for item in toc:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        set_font(p.add_run(item), 10, False, INK)

    doc.add_heading("1. 项目定位与版本范围", level=1)
    add_p(doc, "Kait 是一款把滑行战术盘与 2048 威胁盘放在同一屏幕上的回合制构筑 Demo。玩家每次输入既改变角色在战术盘上的位置，也推动右侧数字盘；数字合成既是资源与构筑来源，也会通过裂隙把新的敌人压力送回战场。")
    add_callout(doc, "当前版本结论", "两名角色共用同一套双盘骨架，但拥有不同的移动、攻击、资源和牌池。Kait 强调高速滑行与连杀路线，Yummn 强调气资源、拳击连携、残影与元素控制。")
    add_simple_table(doc, ["项目项", "当前实现"], [
        ("引擎与画面", "Unity 6；目标画面 1920×1080；Windows 与 Android 双端"),
        ("战术盘", "内部 7×7，实际可操作区域 5×5；两处固定障碍；边界直接停止"),
        ("威胁盘", "5×5 2048 盘；输入方向同步移动、合成并生成新数字"),
        ("角色", "Kait / Yummn，主界面选择后再开始或继续"),
        ("当前 Boss 门槛", "两名角色均在合成 128 后进入盾骑士 Boss 阶段"),
        ("交付状态", "源码、Windows 包、Android APK、自动化测试与本设计文档"),
    ], widths=[1.35, 5.5])

    doc.add_heading("2. 核心体验与完整循环", level=1)
    doc.add_heading("2.1 单次行动循环", level=2)
    add_bullets(doc, [
        "玩家输入方向、等待，或拖出主动技能并选择合法目标。",
        "左侧先结算角色移动、拳击/斩击、推动、击杀、残影和技能即时效果。",
        "右侧威胁盘按规则移动与合成；合法移动后按当前设置生成新的 2。",
        "合成会建立奖励与裂隙队列。裂隙必须先完整预警，之后才刷新敌人。",
        "满足推进条件时进入敌方阶段：敌人统一锁定、统一移动或攻击，并处理友伤、格挡、冻结、震慑等。",
        "回到玩家阶段，更新冷却、气、持续状态、预警和 UI。",
    ])
    doc.add_heading("2.2 长期构筑循环", level=2)
    add_p(doc, "玩家通过合成达到奖励阈值后，从随机牌组中选择新卡，拖到六个装备槽之一。Kait 仍以主动、被动分区组织构筑；Yummn 使用统一六槽，可任意混装主动与被动。奖励界面支持跳过、重抽、替换和收起。")
    add_heading = doc.add_heading("2.3 设计支柱", level=2)
    add_simple_table(doc, ["支柱", "落地方式"], [
        ("一键双解", "同一次方向输入同时解决战场路线与 2048 盘面。"),
        ("读图优先", "攻击和刷新都先给预警；角色站在危险格时加强边缘警示。"),
        ("连贯而不锁手", "移动可打断角色动画；长按方向可连续输入；击杀动画提供自然节拍。"),
        ("跨盘构筑", "卡牌可改变数字移动、补 2、柱子、裂隙与敌人阶段，而不只强化伤害。"),
        ("双角色差异", "Kait 用速度和路线处理敌群；Yummn 用气、拳击与状态组合解决局面。"),
    ], widths=[1.35, 5.5])

    doc.add_heading("3. 双棋盘规则与全局时序", level=1)
    doc.add_heading("3.1 战术盘", level=2)
    add_bullets(doc, [
        "5×5 有效区域内同时容纳角色、敌人、障碍、裂隙、冰柱、暗幕、诱饵与残影。",
        "角色与敌人完整占一格；表现层可越界绘制，但逻辑位置始终按格计算。",
        "裂隙不阻挡敌人移动；若刷新格被占用，则本次刷新取消，不允许无预警直接出怪。",
        "攻击预警和刷新预警位于人物下层；实际命中反馈位于受击单位上层、主角下层。",
    ])
    doc.add_heading("3.2 2048 威胁盘", level=2)
    add_bullets(doc, [
        "滑块使用起步快、收尾慢的曲线，帮助玩家看清方向和合成结果。",
        "默认每次合法输入生成一个 2；Yummn 可通过设置与技能改变补 2 时机和数量。",
        "无可移动方向时立即判负，避免盘面清空或停在不可继续状态。",
        "合成 128 触发当前 Boss 阶段；数值保留为可配置规则，不依赖界面文案硬编码。",
    ])
    doc.add_heading("3.3 表现与逻辑解耦", level=2)
    add_p(doc, "所有单位先完成逻辑快照，再并行播放位移、受击、死亡与攻击动画。敌人死亡动画与 Kait/Yummn 的动画状态机解耦；输入不等待非关键动画结束。敌人死亡时先白闪，再与攻击同步播放 die，原地淡出，显示层位于玩家角色之下。")

    doc.add_heading("4. 角色规则：Kait", level=1)
    add_p(doc, "Kait 是拿大剑的猫娘，核心是沿输入方向高速滑行。碰到敌人时不能穿越：命中并击杀后可进入其位置并继续连杀；未击杀、柱子或边界会终止当前路线。")
    add_simple_table(doc, ["系统", "当前规则"], [
        ("生命", "基础 3 格；敌人攻击、裂隙伤害和碰撞伤害可由设置分别调整。"),
        ("速度", "决定一次滑行可连续处理的格数；连杀等待转向时可继续选择方向。"),
        ("时停提示", "停留后，左半屏边缘逐渐出现大小与方向不同的时钟并灰化，边角密度更高。"),
        ("攻击动画", "普通斩击、击杀、连杀分别反馈；移动可打断所有 Kait 动画。"),
        ("长按输入", "按住当前方向会在一段动作完成后自动继续下一次输入。"),
        ("补 2", "可选仅有效移动至少一格时补 2；时停等待不补 2。"),
    ], widths=[1.3, 5.55])

    doc.add_heading("5. 角色规则：Yummn", level=1)
    add_p(doc, "Yummn 是以气驱动位移和拳法的近战角色。高速与气竭是两套明确的行动状态：有气时待机使用 multiidle；无气时待机使用 idle，移动改用 walk。")
    add_simple_table(doc, ["系统", "当前规则"], [
        ("气槽", "默认上限 6，可在设置选择 3—9；极简侧显示同色圆点，卡通侧使用高质感气点。"),
        ("高速移动", "按移动距离消耗气或固定消耗 1 气；攻击可设为消耗 0 或 1 气。"),
        ("气竭", "气不足仍能出拳；必须把气恢复到上限才退出。气竭移动每次仍补一个 2。"),
        ("气格挡", "默认开启：有气时受到伤害会清空气、进入气竭并抵消本次伤害，2 倍速播放 skill0 前 50 帧。"),
        ("等待", "不移动但推进一个敌方阶段；移动端长按屏幕同一位置触发。等待也生成一个 2。"),
        ("残影", "有效移动至少一格都会生成残影，包括击杀进格和追身步。基础残影承受一次攻击；暗影斗篷可提升为一回合内无限次。"),
        ("交互", "主动技能先拖到战场中央进入准备，再点方向、自身或目标格；释放失败会立即提示并退出准备态。"),
    ], widths=[1.3, 5.55])
    add_callout(doc, "Yummn 敌方阶段口径", "“有效移动推进敌方回合”属于额外开关；原地击杀和气竭原地攻击仍会让敌人行动。")

    doc.add_heading("6. 敌人与 Boss", level=1)
    add_simple_table(doc, ["点数", "兵种", "生命", "战斗职责与攻击规则"], ENEMIES, widths=[0.55, 1.05, 0.55, 4.7])
    doc.add_heading("6.1 敌方阶段", level=2)
    add_bullets(doc, [
        "所有敌人一起完成锁定，再同步移动或攻击，避免逐个行动拖慢节奏。",
        "锁定动作使用 posing 并持续循环，不会自动被待机覆盖。",
        "弓手和术士采用两回合节奏：第一轮锁定，第二轮攻击。即使当前不同行列，弓手也会先尝试选线。",
        "盾骑士对 Yummn 按整排规则锁定和转向；Kait 保留其原有 Boss 逻辑。",
        "敌人可互相伤害；是否启用友伤由角色对应设置决定。",
    ])

    doc.add_heading("7. 完整技能牌池", level=1)
    add_p(doc, "本节按运行时代码中的目录整理。普通、罕见、稀有分别使用白银、蓝、金作为视觉稀有度；主动与被动依靠不同卡面结构区分。实验牌单独标注且默认关闭。")
    add_skill_section(doc, "7.1 Kait 主动技能（14 张）", "主动技能需要拖出并确认目标，使用后进入独立冷却。", KAIT_ACTIVE)
    add_skill_section(doc, "7.2 Kait 被动技能（25 张，含 1 张关闭的实验牌）", "被动技能装备后持续改变战斗、裂隙或威胁盘规则。", KAIT_PASSIVE)
    add_skill_section(doc, "7.3 Yummn 主动技能（10 张）", "方向技能使用四向按钮；疗伤冥想点选自身；塑造流水、黑暗术与次级幻影点选合法空格。", YUMMN_ACTIVE)
    add_skill_section(doc, "7.4 Yummn 被动技能（28 张）", "Yummn 不再区分主动/被动槽位，全部技能共享六个装备位。", YUMMN_PASSIVE)
    add_callout(doc, "技能总量", "当前文档覆盖 Kait 39 张（14 主动 + 25 被动，其中 1 张实验牌关闭）与 Yummn 38 张（10 主动 + 28 被动），合计 77 张目录条目。", CREAM)

    doc.add_heading("8. 奖励、构筑与装备", level=1)
    add_simple_table(doc, ["规则", "实现"], [
        ("Kait 槽位", "3 个主动槽 + 3 个被动槽。"),
        ("Yummn 槽位", "统一 6 槽，主动和被动可自由组合。"),
        ("Yummn 奖励阈值", "默认合成到 16 获得选牌机会；设置可改为 32。"),
        ("卡牌交互", "拖动卡牌到目标槽替换；移出区域取消；卡牌吸附于屏幕边缘并半隐藏。"),
        ("预览", "点击或悬停展示完整卡面；一段时间无操作或点击别处自动缩回。"),
        ("跨裁切线", "同一逻辑卡片使用卡通/极简两套视觉，由斜切线统一裁切，动画状态保持同步。"),
    ], widths=[1.35, 5.5])

    doc.add_heading("9. 操作与界面", level=1)
    add_simple_table(doc, ["平台", "输入"], [
        ("键盘", "W/A/S/D 移动；R 重新开始；等待按钮执行空过；方向长按连续输入。"),
        ("鼠标", "点击预览卡牌；拖出主动技能；点击方向、自身或目标格确认。"),
        ("Android", "滑动映射方向移动；长按当前位置执行等待；长按滑动方向支持连续输入。"),
    ], widths=[1.25, 5.6])
    add_p(doc, "主界面采用 Kait—斜切菜单—Yummn 三段式构图。人物悬浮，鼠标悬停或点击时放大；选中角色后不会立刻进入游戏，统一通过“开始游戏/继续游戏”进入。教程内容随当前角色切换。")
    add_p(doc, "战斗界面保持两个棋盘同尺寸，并向屏幕中心收拢。移动键位于中部；角色状态只保留回合/行动、生命、速度或气等必要信息。可由图标表达的内容不重复堆叠文字。")

    doc.add_heading("10. 视觉、动画与音频", level=1)
    doc.add_heading("10.1 双画风", level=2)
    add_p(doc, "左侧采用翠绿庭院或冰雪庭院的卡通厚涂风：大色块、粗线条、统一俯视角与落地阴影。右侧保留初版简约 2048 风格。中央斜切线贯穿场景、卡片、控制栏、危险边缘与悬浮菜单；同一功能主体只维护一份状态，两侧仅切换表现。")
    doc.add_heading("10.2 角色动画", level=2)
    add_bullets(doc, [
        "Kait：idle 待机；rungamestart 移动；joy_long_return 撞边停止（0.75 倍速）；skill1/skill0/skill2 分别覆盖小型攻击、大型攻击和其他技能；manajump 胜利；踏影使用 000000_run_jump。",
        "Yummn：有气待机 multiidle、气竭待机 idle、滑行 01_run、气竭移动 walk；攻击类技能使用 multistandby，强化 skill1，回复 joyresult，胜利 00000smile。",
        "敌人：landing 出生、idle 待机、posing 锁定、各自 attack / hit / die。单位动画与主角状态机解耦。",
    ])
    doc.add_heading("10.3 特效与音频", level=2)
    add_bullets(doc, [
        "命中、推动、击杀、连杀、格挡分别使用不同层级的特效；击杀与连杀特效位于人物下层，避免遮挡角色。",
        "连杀逐级增强并伴随轻微屏幕震动；不使用全屏红闪。危险将至时屏幕边缘变红，左右画风分别渲染。",
        "弓箭、法术、冰冻、破冰、震慑、残影命中、气收放与裂隙生成均有独立反馈。",
        "Kait 与各兵种语音使用独立通道；新敌人语音出现时旧通道自动衰减，控制多人同时说话的噪声。",
        "音效保留挥剑、普通命中、格挡、击杀、弓箭、法术与界面反馈；背景音乐循环播放并可由设置统一控制。",
    ])

    doc.add_heading("11. 设置、存档与平台", level=1)
    doc.add_heading("11.1 通用与 Kait 设置", level=2)
    add_bullets(doc, [
        "人物无敌；取消 2048 柱子；取消裂隙伤害；取消敌人友伤；取消碰撞伤害。",
        "Kait 可选“有效移动至少一格才生成 2”；时停不生成。",
        "设置界面采用右侧极简风，人物专属选项只在对应角色被选中时出现。",
    ])
    doc.add_heading("11.2 Yummn 设置", level=2)
    add_bullets(doc, [
        "最大气值：3、4、5、6、7、8、9；默认 6。",
        "移动耗气：逐格消耗或高速移动固定 1 气。",
        "高速攻击耗气：0 或 1；气不足仍可出拳并进入气竭。",
        "击杀回气可调，包含回气为 1 的低恢复方案；也可启用回满。",
        "补 2 时机：每次方向操作、每次有效移动、旧版每次移动、仅击杀等互斥方案。",
        "可选“有效移动推进敌方回合”“攻击推进敌方回合”“奖励阈值 16/32”与初始气格挡。",
    ])
    doc.add_heading("11.3 存档", level=2)
    add_p(doc, "存档记录当前角色、战局、盘面、角色状态、敌人、裂隙、技能装备、冷却、奖励与设置。主界面根据有效存档显示统一的“继续游戏”；重新开始会清理当前战局，但保留玩家设置。")
    doc.add_heading("11.4 构建", level=2)
    add_simple_table(doc, ["平台", "当前交付"], [
        ("Windows", "Build/kait.exe，独立运行，不需要 Unity 编辑器。"),
        ("Android", "Build/kait-v0.8.2.apk；包名 com.kaitprototype.demo；ARM64；最低 Android API 26。"),
    ], widths=[1.25, 5.6])

    doc.add_heading("12. 胜负、平衡与测试", level=1)
    add_simple_table(doc, ["项目", "当前判定"], [
        ("胜利", "完成 128 阶段并击败盾骑士 Boss。胜利/死亡动画播放完后不再被待机覆盖。"),
        ("失败", "生命归零，或 2048 威胁盘不存在任何合法移动。"),
        ("裂隙公平性", "裂隙必须经历预警阶段；气竭、合成或其他补 2 不能跳过预警直接刷新。"),
        ("输入响应", "角色移动可打断动画；连续输入不依赖不必要的动画等待；击杀动作提供辨识时间。"),
        ("最近验证", "v0.8.2 震慑/借机攻击相关编辑器测试 67/67 通过；主动技能拖放和三类目标选择完成运行时截图验证。"),
    ], widths=[1.3, 5.55])
    doc.add_heading("12.1 当前平衡观察点", level=2)
    add_bullets(doc, [
        "Yummn 的最大气值、击杀回气与补 2 模式共同决定气竭出现频率，必须按组合而不是单项评估。",
        "疾风连击会放大伺机而动、借机攻击、火蛇之牙等逐拳效果，是构筑中最重要的乘区之一。",
        "永久冰柱和永久暗幕都消耗 2 气且只能保留最新一个，用位置成本约束长期控制。",
        "双盘同时拥挤时，必须保证裂隙预警、敌人预警和角色危险提示仍能清楚分层。",
        "狂野魔法涌动保留为实验卡但默认关闭，不纳入正式平衡评价。",
    ])
    add_callout(doc, "文档维护原则", "本文件描述 2026 年 9 月 13 日项目中的实际实现。后续改动应同步更新技能目录、设置口径、构建信息与测试结论，避免教程、文档和运行时规则再次分叉。", CREAM)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
